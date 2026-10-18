#!/usr/bin/env python3
"""Erzeugt public/data.json aus der Präsentationsliste (Blatt AMBASSADE).

Aufruf:  python3 tools/build_data.py <Liste.xlsx> <B2B-Code>

- B2C-Preise stehen lesbar in data.json.
- B2B-Nettopreise werden mit dem B2B-Code verschlüsselt (PBKDF2-SHA256 + AES-GCM).
  Der Code selbst steht nirgends im Quelltext; die App entschlüsselt erst nach Eingabe.
"""
import json, os, re, sys, base64, unicodedata, hashlib, datetime
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'public', 'data.json')
ITER = 150_000

GROUP = {
    'Rotweine': 'rouge', 'Weissweine': 'blanc', 'Roseweine': 'rose',
    'Schaumweine': 'mousseux', 'Süssweine': 'doux',
    'Weinbrände': 'spiritueux', 'Obstbrände': 'spiritueux',
    'Alkoholfreie Getränke': 'sans',
}
ROOM = {'Salon Pariser Platz': 'S', 'Salle à Manger': 'M', 'Entrée': 'E'}


def norm_code(code: str) -> str:
    s = unicodedata.normalize('NFD', code)
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Z0-9]', '', s.upper())


def price(v):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip()
    if s in ('a.A.', 'a.Z.'):
        return s
    raise ValueError(f'Unbekannter Preiswert: {v!r}')


def liters(u: str) -> float:
    return float(u.replace('l', '').replace(',', '.').strip())


def main(xlsx, code):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb['AMBASSADE']
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    wines, b2b, seen = [], {}, set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[col['WEIN-NR']]:
            continue
        n = str(r[col['WEIN-NR']]).strip()
        if n in seen:
            raise SystemExit(f'Doppelte Wein-Nr {n}')
        seen.add(n)
        g = GROUP[r[col['GRUPPE']].strip()]
        w = {
            'n': n,
            'a': str(r[col['ART. NR.']]).strip(),
            't': int(r[col['TISCH-NR']]),
            'k': g,
            'r': str(r[col['GEBIET']]).strip(),
            'd': str(r[col['DOMAINE']]).strip(),
            'w': str(r[col['WEINBEZEICHNUNG']]).strip(),
            'u': str(r[col['EINHEIT']]).strip(),
            'l': liters(str(r[col['EINHEIT']])),
            'v': (r[col['REBSORTE']] or '').strip(),
            'o': (r[col['BIO-Code']] or '') == 'BIO',
            'm': ROOM[r[col['RAUM']].strip()],
            'c': price(r[col['B2C Brutto']]),
        }
        wines.append(w)
        b2b[n] = price(r[col['B2B Netto']])

    salt, iv = os.urandom(16), os.urandom(12)
    kdf = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER)
    key = kdf.derive(norm_code(code).encode())
    ct = AESGCM(key).encrypt(iv, json.dumps(b2b, separators=(',', ':')).encode(), None)
    b64 = lambda b: base64.b64encode(b).decode()

    payload = {
        'version': datetime.datetime.now().strftime('%Y%m%d-%H%M'),
        'source': os.path.basename(xlsx),
        'count': len(wines),
        'wines': wines,
        'b2b': {'iter': ITER, 'salt': b64(salt), 'iv': b64(iv), 'ct': b64(ct)},
    }
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, separators=(',', ':'))
    print(f'{len(wines)} Positionen -> {OUT} ({os.path.getsize(OUT)//1024} KB), Version {payload["version"]}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
